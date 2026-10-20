"""Export of approved documents as CSV or Excel.

A document counts as approved when its latest review is approve/correct, or when it was
auto-approved and nobody reviewed it since. Rejected documents are left out. Corrected
documents export the values the reviewer saved, not the model output.

Two tables come out: one row per document (header fields flattened to dotted keys such as
`seller.npwp`) and one row per line item, joined by `document_id`.
"""
from __future__ import annotations

import csv
import io
from typing import Any, Iterable, Optional

from sqlalchemy import select
from sqlalchemy.orm import Session

from docex.db.models import Document

META_COLS = ["document_id", "filename", "doc_type", "decision", "reviewer", "decided_at", "approach", "score",
             "received_at"]
ITEM_COLS = ["document_id", "filename", "doc_type", "line_no", "description", "quantity", "unit_price",
             "discount", "total", "is_adjustment"]
SKIP_KEYS = {"doc_type", "line_items", "document_inconsistent"}


def _flatten(values: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    out: dict[str, Any] = {}
    for k, v in values.items():
        if not prefix and k in SKIP_KEYS:
            continue
        if isinstance(v, dict):
            out.update(_flatten(v, f"{prefix}{k}."))
        else:
            out[f"{prefix}{k}"] = v
    return out


def approved(s: Session, doc_type: Optional[str] = None) -> list[dict[str, Any]]:
    """Approved documents, oldest first, each as {meta, values, items}."""
    rows = []
    docs = s.scalars(select(Document).where(Document.status == "done").order_by(Document.created_at)).all()
    for d in docs:
        ex = d.latest
        if ex is None or (doc_type and ex.doc_type != doc_type):
            continue
        review = d.reviews[-1] if d.reviews else None
        if review is not None:
            if review.action == "reject":
                continue
            decision = "corrected" if review.action == "correct" else "approved"
            values = review.corrected_values if review.action == "correct" else ex.result.get("values")
            reviewer, decided_at = review.reviewer, review.created_at
        elif ex.route == "auto_approved":
            decision, values, reviewer, decided_at = "auto_approved", ex.result.get("values"), "", ex.created_at
        else:
            continue  # still waiting in the review queue
        values = values or {}
        meta = {"document_id": d.id, "filename": d.filename, "doc_type": ex.doc_type, "decision": decision,
                "reviewer": reviewer, "decided_at": decided_at.isoformat() if decided_at else "",
                "approach": ex.approach, "score": round(ex.score, 4), "received_at": d.created_at.isoformat()}
        rows.append({"meta": meta, "values": _flatten(values), "items": values.get("line_items") or []})
    return rows


def tables(rows: list[dict[str, Any]]) -> tuple[list[str], list[list[Any]], list[list[Any]]]:
    """(document columns, document rows, item rows). Field columns follow first appearance."""
    fields: list[str] = []
    for r in rows:
        fields += [k for k in r["values"] if k not in fields]
    cols = META_COLS + fields
    doc_rows = [[r["meta"][c] for c in META_COLS] + [r["values"].get(f) for f in fields] for r in rows]
    item_rows = []
    for r in rows:
        m = r["meta"]
        for i, it in enumerate(r["items"], 1):
            item_rows.append([m["document_id"], m["filename"], m["doc_type"], i] +
                             [it.get(c) for c in ITEM_COLS[4:]])
    return cols, doc_rows, item_rows


def to_csv(header: list[str], rows: Iterable[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows([["" if v is None else v for v in r] for r in rows])
    # BOM so Excel opens UTF-8 (names, addresses) correctly
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(cols: list[str], doc_rows: list[list[Any]], item_rows: list[list[Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    sheets = [("Dokumen", cols, doc_rows), ("Rincian barang", ITEM_COLS, item_rows)]
    for i, (title, header, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 48)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
