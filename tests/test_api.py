import io
import time

import pytest
from fastapi.testclient import TestClient
from PIL import Image


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'test.db'}")
    monkeypatch.setenv("DOCEX_STORAGE_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("DOCEX_APPROACH", "A")
    monkeypatch.setenv("QUEUE_BACKEND", "inline")
    import importlib

    import docex.api.app as app_mod
    import docex.db.models as models

    models._Session = None
    importlib.reload(app_mod)
    with TestClient(app_mod.app) as c:
        yield c


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (400, 300), "white").save(buf, format="PNG")
    return buf.getvalue()


def _wait(client, doc_id):
    for _ in range(120):
        d = client.get(f"/documents/{doc_id}").json()
        if d["status"] in ("done", "failed"):
            return d
        time.sleep(0.25)
    raise AssertionError("timeout")


def test_upload_process_review_metrics(client):
    r = client.post("/documents", files={"file": ("x.png", _png(), "image/png")}, data={"doc_type": "receipt"})
    assert r.status_code == 202
    doc_id = r.json()["job_id"]
    d = _wait(client, doc_id)
    assert d["status"] == "done" and d["route"] == "review"
    assert client.get(f"/documents/{doc_id}/preview").headers["content-type"] == "image/png"
    queued = next(q for q in client.get("/review/queue").json() if q["id"] == doc_id)
    assert isinstance(queued["weak_fields"], list)
    r = client.post(f"/review/{doc_id}", json={"action": "correct", "values": {"store_name": "Toko A", "total": "Rp 12.500"}})
    assert r.status_code == 200 and "store_name" in r.json()["changed_fields"]
    assert not any(q["id"] == doc_id for q in client.get("/review/queue").json())
    d = client.get(f"/documents/{doc_id}").json()
    assert d["review"]["corrected_values"]["total"] == 12500.0
    m = client.get("/metrics").json()
    assert m["documents"] == 1 and m["reviews"] == {"correct": 1}
    assert m["corrections_by_field"]["documents"] == 1
    assert {"path": "store_name", "count": 1} in m["corrections_by_field"]["fields"]


def test_rejects_bad_type_and_size(client):
    r = client.post("/documents", files={"file": ("x.txt", b"hello world", "text/plain")})
    assert r.status_code == 415
    big = b"%PDF-" + b"0" * (11 * 1024 * 1024)
    assert client.post("/documents", files={"file": ("x.pdf", big, "application/pdf")}).status_code == 413


def test_llm_approach_unavailable_without_key(client, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    r = client.post("/documents", files={"file": ("x.png", _png(), "image/png")}, data={"approach": "C"})
    assert r.status_code == 400


def test_ui_and_health_for_ui(client):
    r = client.get("/")
    assert r.status_code == 200 and "meja cek dokumen" in r.text
    assert client.get("/web/app.js").status_code == 200
    h = client.get("/health").json()
    assert h["available_approaches"][0] == "A"
    assert "seller.npwp" in h["critical_fields"]["faktur_pajak"]
    samples = client.get("/samples").json()
    if samples:
        img = client.get(f"/samples/{samples[0]['name']}/image")
        assert img.headers["content-type"] == "image/jpeg"


def test_export_approved_documents(client):
    ids = []
    for _ in range(3):
        r = client.post("/documents", files={"file": ("x.png", _png(), "image/png")}, data={"doc_type": "receipt"})
        ids.append(r.json()["job_id"])
    for i in ids:
        _wait(client, i)
    corrected, rejected, waiting = ids
    items = [{"description": "Kopi susu", "quantity": "2", "unit_price": "18.000", "total": "36.000"}]
    client.post(f"/review/{corrected}", json={"action": "correct", "reviewer": "nash",
                                               "values": {"store_name": "Toko A", "total": "36.000", "line_items": items}})
    client.post(f"/review/{rejected}", json={"action": "reject"})

    c = client.get("/export/approved/count").json()
    assert c["documents"] == 1 and c["line_items"] == 1 and c["by_decision"] == {"corrected": 1}

    r = client.get("/export/approved?format=csv")
    assert r.headers["content-type"].startswith("text/csv") and "attachment" in r.headers["content-disposition"]
    lines = r.content.decode("utf-8-sig").splitlines()
    assert len(lines) == 2 and corrected in lines[1] and "Toko A" in lines[1] and "corrected" in lines[1]
    assert waiting not in r.text and rejected not in r.text
    items_csv = client.get("/export/approved?format=csv&part=items").content.decode("utf-8-sig").splitlines()
    assert len(items_csv) == 2 and "Kopi susu" in items_csv[1]

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(client.get("/export/approved?format=xlsx").content))
    assert wb.sheetnames == ["Dokumen", "Rincian barang"]
    assert wb["Dokumen"].max_row == 2 and wb["Rincian barang"]["E2"].value == "Kopi susu"
    assert client.get("/export/approved/count?doc_type=invoice").json()["documents"] == 0


def test_key_alone_does_not_enable_paid_approaches(client, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test")
    monkeypatch.setenv("GEMINI_API_KEY", "test")
    monkeypatch.delenv("DOCEX_ENABLE_LLM", raising=False)
    assert client.get("/health").json()["available_approaches"] == ["A"]
    r = client.post("/documents", files={"file": ("x.png", _png(), "image/png")}, data={"approach": "C"})
    assert r.status_code == 400


def test_weak_fields_lists_critical_fields_weakest_first():
    from docex.api.app import weak_fields

    ext = {"store_name": {"value": "Toko", "confidence": 0.95}, "date": {"value": "2025-01-01", "confidence": 0.62},
           "total": {"value": None, "confidence": 0.0}}
    assert weak_fields("receipt", ext) == [{"path": "total", "confidence": None}, {"path": "date", "confidence": 0.62}]
    ext["line_items"] = [{"total": {"value": 5.0, "confidence": 0.9}}, {"total": {"value": 6.0, "confidence": 0.7}}]
    assert weak_fields("receipt", ext)[2] == {"path": "line_items", "confidence": 0.7}
    assert weak_fields("receipt", None) == [{"path": p, "confidence": None} for p in ("store_name", "date", "total")]


def test_corrections_by_field_counts_documents_and_collapses_line_items():
    from docex.api.app import corrections_by_field

    r = corrections_by_field([["total", "line_items.0.total", "line_items.1.total"], ["total"], ["date"], []])
    assert r == {"documents": 4, "fields": [{"path": "total", "count": 2}, {"path": "date", "count": 1},
                                            {"path": "line_items", "count": 1}]}
